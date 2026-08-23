# services.py
from sqlalchemy.orm import Session
from fastapi.responses import StreamingResponse
from .models import (
    FHistoriaTrabajoDecoracionCabecera,
    FHistoriaTrabajoDecoracionColores,
    FHistoriaTrabajoDecoracionEstaciones,
    FHistoriaTrabajoDecoracionTemperatura
)
from .schemas import CabeceraCrear
from docx import Document
import io
import os

#CREAR
def crear_historia_trabajo(db: Session, datos: CabeceraCrear):
    try:
        nueva_cabecera = FHistoriaTrabajoDecoracionCabecera(
            fecha_elaboracion     = datos.fecha_elaboracion,
            molde                 = datos.molde,
            no_plano              = datos.no_plano,
            identificador_maquina = datos.identificador_maquina,
            tipo_decorado         = datos.tipo_decorado,
            no_archa              = datos.no_archa,
            cant_ranuras          = datos.cant_ranuras,
            no_tazas              = datos.no_tazas,
            velocidad_bpm         = datos.velocidad_bpm,
            tiempo_malla_seg      = datos.tiempo_malla_seg,
            problemas             = datos.problemas,
            observaciones         = datos.observaciones,
            revisado_por          = datos.revisado_por
        )
        db.add(nueva_cabecera)
        db.flush()

        for color in datos.colores:
            db.add(FHistoriaTrabajoDecoracionColores(
                id_htd_cabecera = nueva_cabecera.id_htd_cabecera,
                numero_color    = color.numero_color,
                consumo_pintura = color.consumo_pintura,
                codigo_pintura  = color.codigo_pintura
            ))

        for estacion in datos.estaciones:
            db.add(FHistoriaTrabajoDecoracionEstaciones(
                id_htd_cabecera          = nueva_cabecera.id_htd_cabecera,
                numero_estacion          = estacion.numero_estacion,
                cuerpo_presion_squeegees = estacion.cuerpo_presion_squeegees,
                cuerpo_nro_marco         = estacion.cuerpo_nro_marco,
                cuerpo_tipo_malla        = estacion.cuerpo_tipo_malla,
                cuello_presion_squeegees = estacion.cuello_presion_squeegees,
                cuello_nro_marco         = estacion.cuello_nro_marco,
                cuello_tipo_malla        = estacion.cuello_tipo_malla
            ))

        for zona in datos.temperatura:
            db.add(FHistoriaTrabajoDecoracionTemperatura(
                id_htd_cabecera = nueva_cabecera.id_htd_cabecera,
                numero_zona     = zona.numero_zona,
                set_point       = zona.set_point,
                lado_izq        = zona.lado_izq,
                lado_der        = zona.lado_der
            ))

        db.commit()
        db.refresh(nueva_cabecera)
        return {"status": "Registro Creado con Exito", "id": nueva_cabecera.id_htd_cabecera}

    except Exception as e:
        db.rollback()
        return {"status": "error", "message": str(e)}


#LISTAR
def listar_historias_trabajo(db: Session):
    try:
        registros = db.query(FHistoriaTrabajoDecoracionCabecera).order_by(
            FHistoriaTrabajoDecoracionCabecera.id_htd_cabecera.desc()
        ).all()
        data = [
            {
                "id_htd_cabecera":       r.id_htd_cabecera,
                "fecha_elaboracion":     r.fecha_elaboracion,
                "molde":                 r.molde,
                "no_plano":              r.no_plano,
                "identificador_maquina": r.identificador_maquina,
                "tipo_decorado":         r.tipo_decorado,
                "no_archa":              r.no_archa,
                "cant_ranuras":          r.cant_ranuras,
                "no_tazas":              r.no_tazas,
                "revisado_por":          r.revisado_por
            } for r in registros
        ]
        return {"status": "success", "data": data}
    except Exception as e:
        return {"status": "error", "message": str(e)}


#OBTENER
def obtener_historia_trabajo(db: Session, id_htd_cabecera: int):
    try:
        registro = db.query(FHistoriaTrabajoDecoracionCabecera).filter(
            FHistoriaTrabajoDecoracionCabecera.id_htd_cabecera == id_htd_cabecera
        ).first()
        if not registro:
            return {"status": "error", "message": "Registro no encontrado"}

        data = {
            "id_htd_cabecera":       registro.id_htd_cabecera,
            "fecha_elaboracion":     registro.fecha_elaboracion,
            "molde":                 registro.molde,
            "no_plano":              registro.no_plano,
            "identificador_maquina": registro.identificador_maquina,
            "tipo_decorado":         registro.tipo_decorado,
            "no_archa":              registro.no_archa,
            "cant_ranuras":          registro.cant_ranuras,
            "no_tazas":              registro.no_tazas,
            "velocidad_bpm":         registro.velocidad_bpm,
            "tiempo_malla_seg":      registro.tiempo_malla_seg,
            "problemas":             registro.problemas,
            "observaciones":         registro.observaciones,
            "revisado_por":          registro.revisado_por,
            "colores": [
                {
                    "id_htd_color":    c.id_htd_color,
                    "numero_color":    c.numero_color,
                    "consumo_pintura": c.consumo_pintura,
                    "codigo_pintura":  c.codigo_pintura
                } for c in sorted(registro.colores, key=lambda x: x.numero_color)
            ],
            "estaciones": [
                {
                    "id_htd_estacion":          e.id_htd_estacion,
                    "numero_estacion":          e.numero_estacion,
                    "cuerpo_presion_squeegees": e.cuerpo_presion_squeegees,
                    "cuerpo_nro_marco":         e.cuerpo_nro_marco,
                    "cuerpo_tipo_malla":        e.cuerpo_tipo_malla,
                    "cuello_presion_squeegees": e.cuello_presion_squeegees,
                    "cuello_nro_marco":         e.cuello_nro_marco,
                    "cuello_tipo_malla":        e.cuello_tipo_malla
                } for e in sorted(registro.estaciones, key=lambda x: x.numero_estacion)
            ],
            "temperatura": [
                {
                    "id_htd_temperatura": t.id_htd_temperatura,
                    "numero_zona":        t.numero_zona,
                    "set_point":          t.set_point,
                    "lado_izq":           t.lado_izq,
                    "lado_der":           t.lado_der
                } for t in sorted(registro.temperatura, key=lambda x: x.numero_zona)
            ]
        }
        return {"status": "success", "data": data}

    except Exception as e:
        return {"status": "error", "message": str(e)}

#EDITAR
def editar_historia_trabajo(db: Session, id_htd_cabecera: int, datos: CabeceraCrear):
    try:
        cabecera = db.query(FHistoriaTrabajoDecoracionCabecera).filter(
            FHistoriaTrabajoDecoracionCabecera.id_htd_cabecera == id_htd_cabecera
        ).first()
        if not cabecera:
            return {"status": "error", "message": "Registro no encontrado"}

        cabecera.fecha_elaboracion     = datos.fecha_elaboracion
        cabecera.molde                 = datos.molde
        cabecera.no_plano              = datos.no_plano
        cabecera.identificador_maquina = datos.identificador_maquina
        cabecera.tipo_decorado         = datos.tipo_decorado
        cabecera.no_archa              = datos.no_archa
        cabecera.cant_ranuras          = datos.cant_ranuras
        cabecera.no_tazas              = datos.no_tazas
        cabecera.velocidad_bpm         = datos.velocidad_bpm
        cabecera.tiempo_malla_seg      = datos.tiempo_malla_seg
        cabecera.problemas             = datos.problemas
        cabecera.observaciones         = datos.observaciones
        cabecera.revisado_por          = datos.revisado_por

        db.query(FHistoriaTrabajoDecoracionColores).filter(
            FHistoriaTrabajoDecoracionColores.id_htd_cabecera == id_htd_cabecera
        ).delete()

        db.query(FHistoriaTrabajoDecoracionEstaciones).filter(
            FHistoriaTrabajoDecoracionEstaciones.id_htd_cabecera == id_htd_cabecera
        ).delete()

        db.query(FHistoriaTrabajoDecoracionTemperatura).filter(
            FHistoriaTrabajoDecoracionTemperatura.id_htd_cabecera == id_htd_cabecera
        ).delete()

        for color in datos.colores:
            db.add(FHistoriaTrabajoDecoracionColores(
                id_htd_cabecera = id_htd_cabecera,
                numero_color    = color.numero_color,
                consumo_pintura = color.consumo_pintura,
                codigo_pintura  = color.codigo_pintura
            ))

        for estacion in datos.estaciones:
            db.add(FHistoriaTrabajoDecoracionEstaciones(
                id_htd_cabecera          = id_htd_cabecera,
                numero_estacion          = estacion.numero_estacion,
                cuerpo_presion_squeegees = estacion.cuerpo_presion_squeegees,
                cuerpo_nro_marco         = estacion.cuerpo_nro_marco,
                cuerpo_tipo_malla        = estacion.cuerpo_tipo_malla,
                cuello_presion_squeegees = estacion.cuello_presion_squeegees,
                cuello_nro_marco         = estacion.cuello_nro_marco,
                cuello_tipo_malla        = estacion.cuello_tipo_malla
            ))

        for zona in datos.temperatura:
            db.add(FHistoriaTrabajoDecoracionTemperatura(
                id_htd_cabecera = id_htd_cabecera,
                numero_zona     = zona.numero_zona,
                set_point       = zona.set_point,
                lado_izq        = zona.lado_izq,
                lado_der        = zona.lado_der
            ))

        db.commit()
        db.refresh(cabecera)
        return {"status": "Registro Actualizado con Exito", "id": cabecera.id_htd_cabecera}

    except Exception as e:
        db.rollback()
        return {"status": "error", "message": str(e)}

##EXPORTAR
def exportar_historia_trabajo(db: Session, id_htd_cabecera: int):
    try:
        cabecera = db.query(FHistoriaTrabajoDecoracionCabecera).filter(
            FHistoriaTrabajoDecoracionCabecera.id_htd_cabecera == id_htd_cabecera
        ).first()

        if not cabecera:
            return {"status": "error", "message": "Registro no encontrado"}

        colores     = sorted(cabecera.colores,     key=lambda x: x.numero_color)
        estaciones  = sorted(cabecera.estaciones,  key=lambda x: x.numero_estacion)
        temperatura = sorted(cabecera.temperatura, key=lambda x: x.numero_zona)

        # La ruta funciona excelente tal cual la tienes usando os.path.join
        ruta_plantilla = os.path.join("backend", "plantillas", "Formato_de_Historia_de_Trabajo_de_Decoracion.docx")
        doc = Document(ruta_plantilla)

        tablas = doc.tables

        def escribir_celda(celda, valor):
            texto = str(valor) if valor else "N/A"
            for parrafo in celda.paragraphs:
                runs_bold = [r for r in parrafo.runs if r.bold]
                if runs_bold:
                    runs_bold[0].text      = texto
                    runs_bold[0].font.name = "Arial"
                    runs_bold[0].bold      = True
                    for run in runs_bold[1:]:
                        run.text = ""

        def distribuir_texto(texto, limite=85):
            palabras = texto.split()
            partes = []
            linea_actual = ""
            for palabra in palabras:
                if len(linea_actual) + len(palabra) + 1 <= limite:
                    linea_actual += (" " if linea_actual else "") + palabra
                else:
                    partes.append(linea_actual)
                    linea_actual = palabra
            if linea_actual:
                partes.append(linea_actual)
            return partes

        #TABLA 1 - CABECERA
        escribir_celda(tablas[1].rows[0].cells[0], cabecera.fecha_elaboracion)
        escribir_celda(tablas[1].rows[0].cells[2], cabecera.molde)
        escribir_celda(tablas[1].rows[0].cells[5], cabecera.no_plano)
        escribir_celda(tablas[1].rows[1].cells[0], cabecera.identificador_maquina)
        escribir_celda(tablas[1].rows[1].cells[4], cabecera.tipo_decorado)
        escribir_celda(tablas[1].rows[2].cells[0], cabecera.no_archa)
        escribir_celda(tablas[1].rows[2].cells[1], cabecera.cant_ranuras)
        escribir_celda(tablas[1].rows[2].cells[5], cabecera.no_tazas)
        escribir_celda(tablas[1].rows[3].cells[0], cabecera.velocidad_bpm)
        escribir_celda(tablas[1].rows[3].cells[3], cabecera.tiempo_malla_seg)

        #TABLA 2 - COLORES
        for color in colores:
            col_idx = color.numero_color
            escribir_celda(tablas[2].rows[1].cells[col_idx], color.consumo_pintura)
            escribir_celda(tablas[2].rows[2].cells[col_idx], color.codigo_pintura)

        #TABLA 3 - ESTACIONES
        for estacion in estaciones:
            col_idx = estacion.numero_estacion + 1
            escribir_celda(tablas[3].rows[1].cells[col_idx], estacion.cuerpo_presion_squeegees)
            escribir_celda(tablas[3].rows[2].cells[col_idx], estacion.cuello_presion_squeegees)
            escribir_celda(tablas[3].rows[3].cells[col_idx], estacion.cuerpo_nro_marco)
            escribir_celda(tablas[3].rows[4].cells[col_idx], estacion.cuello_nro_marco)
            escribir_celda(tablas[3].rows[5].cells[col_idx], estacion.cuerpo_tipo_malla)
            escribir_celda(tablas[3].rows[6].cells[col_idx], estacion.cuello_tipo_malla)

        #TABLA 4 - TEMPERATURA
        for zona in temperatura:
            col_idx = zona.numero_zona
            escribir_celda(tablas[4].rows[1].cells[col_idx], zona.set_point)
            escribir_celda(tablas[4].rows[2].cells[col_idx], zona.lado_izq)
            escribir_celda(tablas[4].rows[3].cells[col_idx], zona.lado_der)

        #TABLA 5 - TEXTO LIBRE
        
        #FILA 0 - PROBLEMAS (Cabecera)
        texto_problemas  = cabecera.problemas or "N/A"
        partes_problemas = distribuir_texto(texto_problemas)
        p0 = tablas[5].rows[0].cells[0].paragraphs[0]
        if len(p0.runs) > 0:
            p0.runs[0].text = "Problemas presentados durante la corrida: "
        
        val_p0 = partes_problemas[0] if partes_problemas else "N/A"
        if len(p0.runs) > 1:
            p0.runs[1].text      = val_p0
            p0.runs[1].font.name = "Arial"
            p0.runs[1].bold      = True
            for run in p0.runs[2:]:
                run.text = ""
        else:
            r = p0.add_run(val_p0)
            r.font.name = "Arial"
            r.bold      = True

        #FILAS 1 Y 2 - CONTINUACIÓN DE PROBLEMAS
        for i, fila_idx in enumerate([1, 2]):
            p = tablas[5].rows[fila_idx].cells[0].paragraphs[0]
            texto_fila = partes_problemas[i + 1] if (i + 1) < len(partes_problemas) else ""
            if len(p.runs) > 0:
                p.runs[0].text      = texto_fila
                p.runs[0].font.name = "Arial"
                p.runs[0].bold      = True
                for run in p.runs[1:]:
                    run.text = ""
            elif texto_fila:
                r = p.add_run(texto_fila)
                r.font.name = "Arial"
                r.bold      = True

        #FILA 4 - OBSERVACIONES
        p4 = tablas[5].rows[4].cells[0].paragraphs[0]
        if len(p4.runs) > 0:
            p4.runs[0].text = "Observaciones: "
            
        val_obs = cabecera.observaciones or "N/A"
        if len(p4.runs) > 1:
            p4.runs[1].text      = val_obs
            p4.runs[1].font.name = "Arial"
            p4.runs[1].bold      = True
            for run in p4.runs[2:]:
                run.text = ""
        else:
            r = p4.add_run(val_obs)
            r.font.name = "Arial"
            r.bold      = True

        #FILA 8 - REVISADO POR
        p8 = tablas[5].rows[8].cells[0].paragraphs[0]
        if len(p8.runs) > 0:
            p8.runs[0].text = "Revisado Por: "
            
        val_rev = cabecera.revisado_por or "N/A"
        if len(p8.runs) > 1:
            p8.runs[1].text      = val_rev
            p8.runs[1].font.name = "Arial"
            p8.runs[1].bold      = True
            for run in p8.runs[2:]:
                run.text = ""
        else:
            r = p8.add_run(val_rev)
            r.font.name = "Arial"
            r.bold      = True

        #GENERAR ARCHIVO EN MEMORIA
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        nombre_archivo = f"Formato_de_Historia_de_Trabajo_de_Decoracion_{cabecera.fecha_elaboracion}.docx"

        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f"attachment; filename={nombre_archivo}"}
        )

    except Exception as e:
        return {"status": "error", "message": str(e)}