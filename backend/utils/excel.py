# backend/utils/excel.py
import math
from decimal import Decimal
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO


class EstilosExcel:
    #FUENTE DE CABECERA
    fuente_cabecera = Font(
        name="Arial",
        bold=True,
        color="FFFFFF",
        size=12
    )

    #RELLENO AZUL OSCURO PARA CABECERA
    relleno_cabecera = PatternFill(
        fill_type="solid",
        fgColor="1F4E79"
    )

    #ALINEACION CENTRADA
    alineacion_centro = Alignment(
        horizontal="center",
        vertical="center"
    )

    #BORDE FINO EN TODAS LAS CELDAS
    borde_fino = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    #FUENTE NORMAL PARA DATOS
    fuente_datos = Font(
        name="Arial",
        size=10
    )


def escribir_cabecera(hoja, columnas: list[str], titulo: str = None, fila_inicio: int = 1):
    if titulo:
        hoja.merge_cells(start_row=fila_inicio, start_column=1, end_row=fila_inicio, end_column=len(columnas) + 1)
        celda_titulo = hoja.cell(row=fila_inicio, column=1, value=titulo)
        celda_titulo.font = Font(name="Arial", bold=True, size=14, color="1F4E79")
        celda_titulo.alignment = Alignment(horizontal="center", vertical="center")
        hoja.row_dimensions[fila_inicio].height = 25
        fila_inicio += 1

    todas = ["N°"] + columnas
    max_lineas = 1

    for indice, titulo_col in enumerate(todas, start=1):
        celda = hoja.cell(row=fila_inicio, column=indice, value=titulo_col)
        celda.font      = EstilosExcel.fuente_cabecera
        celda.fill      = EstilosExcel.relleno_cabecera
        celda.alignment = EstilosExcel.alineacion_centro
        celda.border    = EstilosExcel.borde_fino

        letra = get_column_letter(indice)
        
        if indice == 1:
            hoja.column_dimensions[letra].width = 10
            ancho = 10
        else:
            hoja.column_dimensions[letra].width = max(len(titulo_col) + 6, 14)
            ancho = hoja.column_dimensions[letra].width

        # Calculamos si el titulo es muy largo para el ancho de la columna
        largo_titulo = len(str(titulo_col))
        lineas = math.ceil(largo_titulo / ancho)
        if lineas > max_lineas:
            max_lineas = lineas

    # Ajustamos el alto de la fila de la cabecera (mínimo 24)
    alto_cabecera = max(max_lineas * 15, 24)
    hoja.row_dimensions[fila_inicio].height = alto_cabecera

def escribir_filas(hoja, filas: list[dict], columnas: list[str], fila_inicio: int = 2):
    max_lineas_global = 1

    for i, fila in enumerate(filas):
        fila_num = fila_inicio + i
        valor_n = i + 1

        letra_num = get_column_letter(1)
        largo_num = len(str(valor_n))
        ancho_actual_num = hoja.column_dimensions[letra_num].width or 10
        hoja.column_dimensions[letra_num].width = max(ancho_actual_num, largo_num + 4)

        numero = hoja.cell(row=fila_num, column=1, value=valor_n)
        numero.font      = EstilosExcel.fuente_datos
        numero.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        numero.border    = EstilosExcel.borde_fino

        max_lineas_fila = 1

        for col_num, clave in enumerate(columnas, start=2):
            valor = fila.get(clave)
            celda = hoja.cell(row=fila_num, column=col_num, value=valor)
            celda.font      = EstilosExcel.fuente_datos
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            celda.border    = EstilosExcel.borde_fino
            
            if isinstance(valor, (int, float, Decimal)):
                celda.number_format = '#,##0.00'
            
            letra = get_column_letter(col_num)
            texto = str(valor) if valor is not None else ""
            
            largo_texto = len(texto)
            ancho_actual = hoja.column_dimensions[letra].width or 14
            nuevo_ancho = max(ancho_actual, largo_texto + 4)
            hoja.column_dimensions[letra].width = min(nuevo_ancho, 60)

            ancho_final = hoja.column_dimensions[letra].width
            lineas_n = texto.count('\n') + 1
            lineas_wrap = math.ceil(largo_texto / ancho_final) if ancho_final > 0 else 1
            lineas_celda = max(lineas_n, lineas_wrap)
            
            if lineas_celda > max_lineas_fila:
                max_lineas_fila = lineas_celda

        if max_lineas_fila > max_lineas_global:
            max_lineas_global = max_lineas_fila

    # Alto mínimo de 24 para que las filas de datos no se vean delgadas
    alto_uniforme = max(max_lineas_global * 15, 24)

    for i in range(fila_inicio, fila_inicio + len(filas)):
        hoja.row_dimensions[i].height = alto_uniforme

def crear_libro(nombre_hoja: str):
    #CREA EL WORKBOOK Y RENOMBRA LA HOJA ACTIVA
    libro = Workbook()
    hoja  = libro.active
    hoja.title = nombre_hoja
    return libro, hoja


def libro_a_bytes(libro: Workbook) -> bytes:
    #CONVIERTE EL LIBRO A BYTES SIN GUARDAR EN DISCO
    buffer = BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()



def escribir_totales(hoja, columnas: list[str], cantidad_filas: int, fila_inicio: int = 2):
    if cantidad_filas == 0:
        return

    fila_total = fila_inicio + cantidad_filas
    fila_rango_inicio = fila_inicio
    fila_rango_fin = fila_inicio + cantidad_filas - 1

    for col_num, clave in enumerate(columnas, start=2):
        letra = get_column_letter(col_num)
        valor = ""
        formato = "General"

        if "Eficiencia" in clave or "Deficiencia" in clave:
            valor = f"=AVERAGE({letra}{fila_rango_inicio}:{letra}{fila_rango_fin})"
            formato = '#,##0.00"%"'
        elif clave not in ["Fecha", "Turno", "Grupo", "Técnico", "Moldes", "Nombre", "Descripción", "Item"]:
            valor = f"=SUM({letra}{fila_rango_inicio}:{letra}{fila_rango_fin})"
            formato = "#,##0.00"
        elif col_num == 2:
            valor = "TOTALES"

        celda = hoja.cell(row=fila_total, column=col_num, value=valor)
        celda.font      = Font(name="Arial", bold=True, color="FFFFFF", size=12)
        celda.fill      = EstilosExcel.relleno_cabecera 
        celda.alignment = Alignment(horizontal="center", vertical="center")
        celda.border    = EstilosExcel.borde_fino
        celda.number_format = formato

    celda_num = hoja.cell(row=fila_total, column=1)
    celda_num.fill = EstilosExcel.relleno_cabecera
    celda_num.border = EstilosExcel.borde_fino
    hoja.row_dimensions[fila_total].height = 25